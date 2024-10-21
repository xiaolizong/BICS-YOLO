
from docx import Document
import xlwt
import datetime
import numpy as np


class rwtes():

    def __init__(self,file_name,aim_file_name):

        self.file_name = file_name
        self.aim_file_name = aim_file_name

    # 读取word中table
    def read_word_table(self):
        try:
            data_dict = {}  # 用来存储数据
            doc = Document(self.file_name)  # 读取文档
            num = 0
            flag = True
            while flag:
                table_list = []
                try:
                    one_tbale = doc.tables[num]  # 提取一个表格对象
                    tb1 = one_tbale  # 提取一个表格对象

                    for i in range(0, len(tb1.rows)):
                        row_lis = []
                        row = tb1.rows[i]  # 提取表格的每一行对象
                        tb1_row_cells = row.cells  # 提取到行里面的所有的格子
                        for j in range(len(tb1_row_cells)):
                            tb1_row_cell = tb1_row_cells[j]
                            cell_text = tb1_row_cell.text
                            row_lis.append(cell_text)
                        table_list.append(row_lis)
                    data_dict[str(num)] = table_list
                    now3 = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                    print('{} >>> word中表{}读取完成'.format(now3, num+1))
                    num += 1
                except:
                    flag = False
            return data_dict
        except Exception as e:
            print(e)

    def save_table_to_sheet(self, table_list):
        try:
            # 创建excel表格类型文件
            book = xlwt.Workbook(encoding='utf-8', style_compression=0)
            sheet_num = 1
            for value in table_list.values():
                # 列名称设置
                # col_name = value[0]
                table_list_num = np.array(value)   # np_var_sheet2
                rows = table_list_num.shape[0]  # 输出行
                cols = table_list_num.shape[1]  # 输出列
                # 在excel表格类型文件中建立一张sheet表单
                sheet = book.add_sheet('sheet{}'.format(sheet_num), cell_overwrite_ok=True)
                # 数据写入
                for i in range(0,rows):
                    for j in range(0,cols):
                        sheet.write(i, j, value[i][j])
                now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
                print('{}>>> sheet{}已经写入！'.format(now, sheet_num))
                sheet_num += 1
            now2 = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
            savepath = r'C:\****\output\{}{}.xls'.format(self.aim_file_name, now2)
            book.save(savepath)
            now4 = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            print('{}>>> 文件已经保存！！'.format(now4))
        except Exception as e:
            print(e)

if __name__ == "__main__":
    # 给入word的路径
    filename =r"C:\****\Excel_数据质量20230419\附件：数据报告-2023001.docx"
    aim_file_name = r'数据质量表格1'
    rwtes = rwtes(filename, aim_file_name)
    data_dict = rwtes.read_word_table()
    print('>>>>>>>>>>>>>>>>>>>>>>>>>>>>>>>>>>>>>>>>>>>>')
    rwtes.save_table_to_sheet(data_dict)

"""
该程序有3个参数需要修改
para1：filename —— word存储路径
para2：aim_file_name —— excelm名称
para3: save_table_to_sheet函数中excel保存位置
"""
